"""Offline tests for create_pptx's infographic layouts.

cards / steps / timeline / comparison / progress, the split chart|image slide,
the takeaway banner and the icon set. Same harness as test_create_pptx.py: the
tool fn runs directly against a temp-configured fallback store, and the saved
deck is reopened with python-pptx. Kept in its own file so the original
suite's shape stays as it was.
"""

import asyncio

import pytest

from app.files.store import file_store
from app.tools.local import pptx as pptx_tool
from app.tools.local import pptx_icons


@pytest.fixture(autouse=True)
def _configure_store(tmp_path):
    file_store.configure(str(tmp_path))
    yield


def _run(args):
    return asyncio.run(pptx_tool.SPEC.func(args))


def _record(result: str):
    assert "Download it at: GET /v1/files/" in result, result
    record = file_store.get(result.split("/v1/files/")[1].strip().split()[0])
    assert record is not None
    return record


def _slides(result: str):
    from pptx import Presentation

    return list(Presentation(_record(result).path).slides)


def _texts(slide) -> list[str]:
    return [p.text for s in slide.shapes if s.has_text_frame for p in s.text_frame.paragraphs if p.text]


def _pictures(slide) -> int:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    return sum(1 for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)


CARDS = [
    {"icon": "shield-check", "title": "Secure", "text": "Zero critical findings"},
    {"icon": "growth", "title": "Growth", "text": "18% fee income"},
    {"icon": "no-such-icon", "title": "Talent"},
]


# --- each layout renders --------------------------------------------------------


def test_cards_render_titles_text_and_one_picture_per_resolved_icon():
    result = _run({"slides": [{"title": "Pillars", "cards": CARDS}]})
    content = _slides(result)[0]
    texts = _texts(content)
    for expected in ("Pillars", "Secure", "Zero critical findings", "Growth", "Talent"):
        assert expected in texts
    # 'growth' is an alias (-> trending-up); 'no-such-icon' is dropped.
    assert _pictures(content) == 2


@pytest.mark.parametrize(
    "key,items,expected",
    [
        ("steps", [{"title": "Apply", "text": "Online", "icon": "smartphone"}, {"title": "Approve"}], ["Apply", "Online", "Approve"]),
        ("timeline", [{"date": "Q1 2026", "title": "Relaunch"}, {"date": "Q2 2026", "title": "Loans", "text": "Instant"}], ["Q1 2026", "Relaunch", "Q2 2026", "Instant"]),
        ("comparison", [{"heading": "Branch", "points": ["Slow"]}, {"heading": "Digital", "icon": "smartphone", "points": ["Fast", "24/7"]}], ["Branch", "Slow", "Digital", "24/7"]),
        ("progress", [{"label": "Adoption", "value": 72, "note": "Target 80%"}, {"label": "NPS", "value": "64%"}], ["Adoption", "Target 80%", "72%", "64%"]),
    ],
)
def test_each_layout_renders_its_text(key, items, expected):
    result = _run({"slides": [{"title": "T", key: items}]})
    texts = " | ".join(_texts(_slides(result)[0]))
    for text in expected:
        assert text in texts, (text, texts)


def test_takeaway_renders_as_its_own_line():
    result = _run({"slides": [{"title": "T", "cards": CARDS[:2], "takeaway": "One key message"}]})
    assert "One key message" in _texts(_slides(result)[0])


def test_split_chart_slide_renders_chart_and_side_bullets():
    chart = {"chart_type": "bar", "labels": ["Q1", "Q2"], "series": [{"name": "D", "data": [1, 2]}]}
    result = _run({"slides": [{"title": "T", "chart": chart, "bullets": ["Up every quarter"]}]})
    slide = _slides(result)[0]
    assert any(s.has_chart for s in slide.shapes)
    assert any("Up every quarter" in t for t in _texts(slide))


def test_split_chart_slide_renders_stacked_stats():
    chart = {"chart_type": "line", "labels": ["a", "b"], "series": [{"data": [1, 2]}]}
    stats = [{"value": "+24%", "label": "YoY"}, {"value": "262B", "label": "Deposits"}]
    result = _run({"slides": [{"title": "T", "chart": chart, "stats": stats}]})
    texts = _texts(_slides(result)[0])
    assert "+24%" in texts and "262B" in texts


def test_single_series_chart_has_no_legend_and_multi_series_does():
    one = {"chart_type": "bar", "labels": ["a", "b"], "series": [{"name": "s", "data": [1, 2]}]}
    two = {"chart_type": "bar", "labels": ["a", "b"], "series": [{"name": "x", "data": [1, 2]}, {"name": "y", "data": [2, 1]}]}
    result = _run({"slides": [{"title": "1", "chart": one}, {"title": "2", "chart": two}]})
    charts = [next(s.chart for s in slide.shapes if s.has_chart) for slide in _slides(result)[:2]]
    assert charts[0].has_legend is False
    assert charts[1].has_legend is True


# --- normalization and the preview --------------------------------------------------


def test_preview_carries_the_resolved_icons_not_the_raw_names():
    result = _run({"slides": [{"title": "T", "cards": CARDS}]})
    cards = _record(result).preview["slides"][0]["cards"]
    assert [c["icon"] for c in cards] == ["shield-check", "trending-up", None]


def test_progress_values_are_numeric_in_the_preview():
    result = _run({"slides": [{"title": "T", "progress": [{"label": "A", "value": "72%"}]}]})
    assert _record(result).preview["slides"][0]["progress"][0]["value"] == 72.0


def test_resolve_icon_aliases_spacing_and_unknowns():
    assert pptx_icons.resolve_icon("Shield Check") == "shield-check"
    assert pptx_icons.resolve_icon("money") == "banknote"
    assert pptx_icons.resolve_icon("definitely-not-an-icon") is None
    assert pptx_icons.resolve_icon(None) is None
    assert all(target in pptx_icons.ICON_NAMES for target in pptx_icons.ALIASES.values())


def test_every_icon_name_has_both_committed_assets():
    """ICON_NAMES and the PNGs drift apart silently otherwise: a missing asset
    renders an icon-less card with no error (scripts/build_pptx_icons.py)."""
    for name in pptx_icons.ICON_NAMES:
        for variant in ("red", "white"):
            assert pptx_icons.icon_path(name, variant) is not None, (name, variant)
    assert (pptx_icons.ICON_DIR / "LICENSE").exists()


# --- refusals -----------------------------------------------------------------------


# --- over-packed slides are REPAIRED, not refused --------------------------------
# Measured live: qwen3.5:35b-a3b puts bullets beside cards on most slides, and
# refusing that ended the turn at max_iterations with no deck (_split_overpacked).


def _layouts(result: str) -> list[list[str]]:
    keys = ("cards", "steps", "timeline", "comparison", "progress", "stats", "chart", "image", "table", "bullets", "takeaway")
    return [[k for k in keys if s.get(k) is not None] for s in _record(result).preview["slides"]]


def test_two_layouts_on_one_slide_become_two_slides():
    result = _run({"slides": [{"title": "T", "cards": CARDS[:2], "steps": [{"title": "a"}, {"title": "b"}]}]})
    assert _layouts(result) == [["cards"], ["steps"]]
    titles = [s["title"] for s in _record(result).preview["slides"]]
    assert titles == ["T", "T"]


@pytest.mark.parametrize(
    "extra,moved",
    [
        ({"bullets": ["x"]}, ["bullets"]),
        ({"stats": [{"value": "1", "label": "l"}]}, ["stats"]),
        ({"table": {"rows": [["a"]]}}, ["table"]),
    ],
)
def test_extras_beside_a_layout_move_to_a_continuation_slide(extra, moved):
    result = _run({"slides": [{"title": "T", "cards": CARDS[:2], "takeaway": "msg", **extra}]})
    assert _layouts(result) == [["cards", "takeaway"], moved]
    slides = _slides(result)
    assert "x" in _texts(slides[1]) or "1" in _texts(slides[1]) or any(s.has_table for s in slides[1].shapes)


def test_a_side_panel_with_both_bullets_and_stats_keeps_the_stats_beside_the_chart():
    chart = {"chart_type": "bar", "labels": ["a"], "series": [{"data": [1]}]}
    result = _run({"slides": [{"title": "T", "chart": chart, "bullets": ["a"], "stats": [{"value": "1", "label": "l"}]}]})
    assert _layouts(result) == [["stats", "chart"], ["bullets"]]


def test_too_many_side_bullets_move_to_their_own_slide():
    chart = {"chart_type": "bar", "labels": ["a"], "series": [{"data": [1]}]}
    bullets = [f"b{i}" for i in range(pptx_tool.MAX_SIDE_BULLETS + 1)]
    result = _run({"slides": [{"title": "T", "chart": chart, "bullets": bullets}]})
    assert _layouts(result) == [["chart"], ["bullets"]]


def test_splitting_still_respects_the_slide_cap():
    over = [{"title": "T", "cards": CARDS[:2], "bullets": ["x"]}] * (pptx_tool.MAX_SLIDES // 2 + 1)
    result = _run({"slides": over})
    assert result.startswith("ERROR:") and "split" in result


def test_takeaway_with_bullets_renders_and_with_a_table_is_refused():
    ok = _run({"slides": [{"title": "T", "bullets": ["a", "b"], "takeaway": "msg"}]})
    assert "msg" in _texts(_slides(ok)[0])
    refused = _run({"slides": [{"title": "T", "table": {"rows": [["a"]]}, "takeaway": "msg"}]})
    assert refused.startswith("ERROR:") and "takeaway" in refused


@pytest.mark.parametrize(
    "key,items",
    [
        ("cards", []),
        ("cards", [{"title": f"c{i}"} for i in range(pptx_tool.MAX_CARDS + 1)]),
        ("comparison", [{"heading": f"h{i}", "points": ["p"]} for i in range(pptx_tool.MAX_COLUMNS + 1)]),
        ("progress", []),
    ],
)
def test_item_counts_are_bounded(key, items):
    result = _run({"slides": [{"title": "T", key: items}]})
    assert result.startswith("ERROR:") and key in result


def test_long_copy_is_set_smaller_not_refused():
    """Measured live: refusing a title a few characters over its limit looped
    the model to max_iterations with no deck. Long copy shrinks instead."""
    comfortable = pptx_tool._TEXT_CAPS["card.text"]
    for length, size in ((comfortable, 12), (comfortable + 1, 10), (int(comfortable * 1.5) + 1, 8)):
        assert pptx_tool._fit("x" * length, 12, "card.text") == size, length
    result = _run({"slides": [{"title": "T", "cards": [{"title": "a", "text": "x" * (comfortable + 20)}, {"title": "b"}]}]})
    assert result.startswith("Created"), result


def test_a_single_item_layout_renders():
    result = _run({"slides": [{"title": "T", "comparison": [{"heading": "Only", "points": ["one"]}]}]})
    assert result.startswith("Created"), result


def test_paragraph_length_copy_is_refused():
    paragraph = "x" * (pptx_tool._TEXT_CAPS["card.text"] * pptx_tool._HARD_CAP_FACTOR + 1)
    result = _run({"slides": [{"title": "T", "cards": [{"title": "a", "text": paragraph}, {"title": "b"}]}]})
    assert result.startswith("ERROR:") and "paragraph" in result


@pytest.mark.parametrize("value", [101, -1, "lots"])
def test_progress_value_must_be_a_percentage(value):
    result = _run({"slides": [{"title": "T", "progress": [{"label": "A", "value": value}]}]})
    assert result.startswith("ERROR:") and "0" in result and "100" in result


def test_timeline_needs_dates():
    result = _run({"slides": [{"title": "T", "timeline": [{"title": "a"}, {"title": "b"}]}]})
    assert result.startswith("ERROR:") and "date" in result


# --- routing --------------------------------------------------------------------------


def test_description_leads_with_the_infographic_rule_and_names_every_layout():
    """The description is the routing prompt: the design rule must come BEFORE
    the field list (CLAUDE.md: a buried hint stops being read), and every
    layout the model may choose must be named."""
    description = pptx_tool.SPEC.description
    assert description.index("DESIGN RULE") < description.index("'cards'")
    for key in ("cards", "steps", "timeline", "comparison", "progress", "stats", "chart", "image", "takeaway"):
        assert f"'{key}'" in description, key
    slide_props = pptx_tool.SPEC.parameters["properties"]["slides"]["items"]["properties"]
    for key in pptx_tool._LAYOUT_KEYS + ("takeaway",):
        assert key in slide_props, key


# --- JSON-encoded strings (measured: the model sends `slides` as a string) -------------


def test_slides_sent_as_a_json_string_are_decoded():
    import json

    slides = [{"title": "T", "cards": CARDS[:2], "takeaway": "msg"}]
    result = _run({"slides": json.dumps(slides)})
    assert _layouts(result) == [["cards", "takeaway"]]


def test_a_layout_field_sent_as_a_json_string_is_decoded():
    import json

    result = _run({"slides": [{"title": "T", "steps": json.dumps([{"title": "a"}, {"title": "b"}])}]})
    assert _layouts(result) == [["steps"]]


def test_a_string_that_is_not_json_is_still_refused_with_guidance():
    result = _run({"slides": "three slides about loans"})
    assert result.startswith("ERROR:") and "not a string" in result
