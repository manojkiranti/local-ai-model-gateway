"""create_memo — the bank's fixed memo format, filled with the model's content.

Pure: the in-memory file store, python-docx to reopen what was written. No
database, no model."""

import asyncio
from datetime import date

import pytest

from app.files.store import DOCX_MEDIA_TYPE, file_store
from app.tools.local import memo as memo_tool


@pytest.fixture(autouse=True)
def _configure_store(tmp_path):
    file_store.configure(str(tmp_path))
    yield


def _run(args):
    return asyncio.run(memo_tool.SPEC.func(args))


def _link_id(result: str) -> str:
    assert "Download it at: GET /v1/files/" in result, result
    return result.split("/v1/files/")[1].strip().split()[0]


def _args(**overrides):
    args = {
        "to": "Chief Executive Officer",
        "from": "Compliance Department",
        "subject": "Approval for Fonepay V2 API Implementation",
        "date": "5th August, 2025",
        "objective": ["This memo seeks **approval** for X."],
        "background": ["FIU Nepal requires contra details."],
        "recommendation": ["Approve the proposal."],
        "sections": [
            {
                "heading": "Risk and Mitigation:",
                "content": [
                    {"heading": "Unauthorized Access"},
                    {"bullets": ["**Risk**: leakage.", {"text": "**Mitigation**:", "bullets": ["Encrypt."]}]},
                ],
            },
        ],
        "signatories": [
            {"role": "Prepared By", "name": "Shristi Bajracharya", "designation": "Assistant Compliance"},
            {"role": "Supported By", "name": "Dipendra Sharma", "designation": "Head Compliance"},
            {"role": "Supported By", "name": "Ranjeet Thakur", "designation": "Officer-DTE"},
            {"role": "Approved By", "name": "Roshan Kumar Neupane", "designation": "Chief Executive Officer"},
        ],
    }
    args.update(overrides)
    return args


def _document(result):
    from docx import Document

    record = file_store.get(_link_id(result))
    assert record.media_type == DOCX_MEDIA_TYPE
    return Document(record.path), record


def _all_text(document) -> str:
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


# ---- validation -------------------------------------------------------------


@pytest.mark.parametrize("key", ["to", "from", "subject"])
def test_to_from_subject_are_required(key):
    assert _run(_args(**{key: " "})).startswith(f"ERROR: '{key}'")


def test_signatories_are_required():
    assert _run(_args(signatories=[])).startswith("ERROR: 'signatories'")


@pytest.mark.parametrize("key", ["objective", "background", "recommendation"])
def test_objective_background_and_recommendation_are_required(key):
    result = _run(_args(**{key: None}))
    assert result.startswith(f"ERROR: every memo needs an '{key}' section"), result


def test_other_sections_are_optional():
    result = _run(_args(sections=None))
    assert result.startswith("Created memo"), result
    _, record = _document(result)
    assert [s["heading"] for s in record.preview["sections"]] == [
        "Objective:",
        "Background:",
        "Recommendation and Conclusion:",
    ]


def test_required_sections_found_inside_sections_are_moved_into_place():
    result = _run(
        _args(
            objective=None,
            background=None,
            recommendation=None,
            sections=[
                {"heading": "Recommendation:", "content": ["Approve."]},
                {"heading": "Cost", "content": ["NPR 3,390."]},
                {"heading": "Background", "content": ["Why."]},
                {"heading": "Objective", "content": ["What."]},
            ],
        )
    )
    _, record = _document(result)
    assert [s["heading"] for s in record.preview["sections"]] == [
        "Objective",
        "Background",
        "Cost",
        "Recommendation:",
    ]


def test_a_bad_block_is_refused_with_its_location():
    result = _run(_args(sections=[{"heading": "X", "content": [{"table": []}]}]))
    assert result.startswith("ERROR: sections[0].content[0]")


def test_a_signatory_needs_a_role_and_a_name():
    result = _run(_args(signatories=[{"role": "Prepared By"}]))
    assert result.startswith("ERROR: signatories[0]")


def test_more_sections_than_letters_is_refused():
    many = [{"heading": f"S{i}", "content": ["x"]} for i in range(memo_tool.MAX_SECTIONS)]
    assert "limit" in _run(_args(sections=many))


# ---- the fixed format -------------------------------------------------------


def test_the_memo_carries_the_fixed_format_and_the_dynamic_content():
    result = _run(_args())
    assert result.startswith("Created memo"), result
    document, _ = _document(result)
    text = _all_text(document)

    # Fixed by the format.
    assert "MEMO" in text
    assert memo_tool.APPROVAL_HEADING in text
    assert len(document.inline_shapes) == 1, "the logo is missing"
    # Dynamic, from the arguments.
    for value in ("Chief Executive Officer", "Compliance Department", "Approval for Fonepay V2"):
        assert value in text
    # Sections are lettered; numbered sub-headings count within a section.
    # Objective, Background, the optional sections, Recommendation — lettered.
    assert "A. Objective:" in text
    assert "B. Background:" in text
    assert "C. Risk and Mitigation:" in text
    assert "D. Recommendation and Conclusion:" in text
    assert "1. Unauthorized Access" in text


def test_header_uses_no_tables_so_no_outline_shows():
    """MEMO + logo and To/From/Subject/Date are paragraphs with tab stops; the
    only table in the memo is the signature grid."""
    document, _ = _document(_run(_args(sections=None)))
    assert len(document.tables) == 1
    lines = [p.text for p in document.paragraphs[:5]]
    assert lines[0].startswith("MEMO")
    assert lines[1:5] == [
        "To\t:\tChief Executive Officer",
        "From\t:\tCompliance Department",
        "Subject\t:\tApproval for Fonepay V2 API Implementation",
        "Date\t:\t5th August, 2025",
    ]


def test_the_header_block_is_ruled_above_to_and_below_date_only():
    from docx.oxml.ns import qn

    document, _ = _document(_run(_args()))
    rules = []
    for p in document.paragraphs[1:5]:
        borders = p._p.pPr.find(qn("w:pBdr"))
        rules.append(sorted(e.tag.split("}")[1] for e in borders) if borders is not None else [])
    assert rules == [["top"], [], [], ["bottom"]]


def test_section_headings_are_bold_not_underlined():
    document, _ = _document(_run(_args()))
    heading = next(p for p in document.paragraphs if p.text.startswith("A. Objective"))
    assert all(r.bold and not r.underline for r in heading.runs)


def test_signatories_fill_a_three_column_grid_in_order():
    document, _ = _document(_run(_args()))
    grid = document.tables[-1]
    # Two groups of three rows: role / signing space / name + designation.
    assert len(grid.rows) == 6
    assert [c.text for c in grid.rows[0].cells] == ["Prepared By", "Supported By", "Supported By"]
    assert grid.rows[2].cells[0].text == "Shristi Bajracharya\nAssistant Compliance"
    assert grid.rows[3].cells[0].text == "Approved By"
    assert grid.rows[5].cells[0].text.startswith("Roshan Kumar Neupane")
    assert grid.rows[3].cells[1].text == ""  # the unused cells stay blank


def test_bold_markup_becomes_bold_runs():
    document, _ = _document(_run(_args()))
    para = next(p for p in document.paragraphs if "seeks" in p.text)
    assert "**" not in para.text
    assert [r.text for r in para.runs if r.bold] == ["approval"]


def test_everything_is_set_in_arial():
    document, _ = _document(_run(_args()))
    runs = [r for p in document.paragraphs for r in p.runs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                runs.extend(r for p in cell.paragraphs for r in p.runs)
    assert {r.font.name for r in runs if r.text.strip()} == {memo_tool.MEMO_FONT}


def test_the_date_defaults_to_today_in_the_memo_style(monkeypatch):
    monkeypatch.setattr(memo_tool, "today", lambda: date(2025, 8, 5))
    document, _ = _document(_run(_args(date=None)))
    assert document.paragraphs[4].text == "Date\t:\t5th August, 2025"


@pytest.mark.parametrize(
    ("day", "expected"),
    [(1, "1st"), (2, "2nd"), (3, "3rd"), (4, "4th"), (11, "11th"), (12, "12th"), (13, "13th"), (21, "21st"), (22, "22nd"), (23, "23rd")],
)
def test_ordinal_suffixes(day, expected):
    assert memo_tool.format_memo_date(date(2025, 8, day)).startswith(f"{expected} August")


def test_saved_record_carries_a_memo_preview():
    _, record = _document(_run(_args()))
    assert record.preview["kind"] == "memo"
    assert record.preview["to"] == "Chief Executive Officer"
    assert record.preview["signatories"][0] == {
        "role": "Prepared By",
        "name": "Shristi Bajracharya",
        "designation": "Assistant Compliance",
    }


def test_filename_suffix_is_forced():
    _, record = _document(_run(_args(filename="fonepay-approval")))
    assert record.filename == "fonepay-approval.docx"


# ---- registration -----------------------------------------------------------


def test_tool_is_registered_once():
    from app.tools.local import LOCAL_TOOLS

    assert [t.name for t in LOCAL_TOOLS].count("create_memo") == 1


def test_description_routes_other_word_documents_to_create_docx():
    assert "create_docx" in memo_tool.SPEC.description


# ---- tolerant input ---------------------------------------------------------
# Shapes the deployed model actually sent (2026-09-30): refused ten times, it
# fell back to create_docx and the user got an unbranded document.


def _sections(result):
    """The optional sections only — between Background and Recommendation."""
    _, record = _document(result)
    return record.preview["sections"][2:-1]


def test_sections_sent_as_a_json_string_are_decoded():
    import json

    result = _run(_args(sections=json.dumps([{"heading": "Objective", "content": ["Text."]}])))
    assert result.startswith("Created memo"), result
    assert _sections(result) == [{"heading": "Objective", "content": ["Text."]}]


def test_a_left_out_closing_bracket_is_repaired():
    # The model's own mistake: the bullet list's "]" left out before "}}".
    broken = (
        '[{"heading": "Background", "content": [{"type": {"bullets": ["a", "b"}}]}, '
        '{"heading": "Next", "content": [{"type": {"text": "c"}}]}]'
    )
    result = _run(_args(sections=broken))
    assert result.startswith("Created memo"), result
    assert _sections(result) == [
        {"heading": "Background", "content": [{"bullets": ["a", "b"]}]},
        {"heading": "Next", "content": ["c"]},
    ]


def test_unrepairable_json_is_refused_with_an_example():
    result = _run(_args(sections='[{"heading": "X", "content": [oops]}]'))
    assert result.startswith("ERROR: 'sections' was sent as a string")
    assert "Example:" in result


@pytest.mark.parametrize(
    ("block", "expected"),
    [
        ({"paragraph": "Hello"}, "Hello"),
        ({"type": {"text": "Hello"}}, "Hello"),
        ({"type": "paragraph", "text": "Hello"}, "Hello"),
        ({"subheading": "Risk"}, {"heading": "Risk"}),
        ({"type": "bullets", "items": ["a"]}, {"bullets": ["a"]}),
        ({"type": {"bullets": ["a"]}}, {"bullets": ["a"]}),
        ("- a\n- b", {"bullets": ["a", "b"]}),
    ],
)
def test_block_shapes_are_normalised(block, expected):
    result = _run(_args(sections=[{"heading": "S", "content": [block]}]))
    assert _sections(result)[0]["content"] == [expected]


def test_dash_lines_inside_a_bullet_become_sub_bullets():
    block = {"bullets": ["Mitigation:\n  - Encrypt\n  - Audit"]}
    result = _run(_args(sections=[{"heading": "S", "content": [block]}]))
    assert _sections(result)[0]["content"] == [
        {"bullets": [{"text": "Mitigation:", "bullets": ["Encrypt", "Audit"]}]}
    ]


def test_a_lead_in_paragraph_with_dash_lines_splits_into_paragraph_and_list():
    result = _run(_args(sections=[{"heading": "S", "content": ["Categories:\n- A class\n- B class"]}]))
    assert _sections(result)[0]["content"] == ["Categories:", {"bullets": ["A class", "B class"]}]


def test_a_letter_in_the_heading_is_not_doubled():
    result = _run(_args(objective=None, sections=[{"heading": "A. Objective", "content": ["x"]}]))
    document, _ = _document(result)
    assert "A. Objective" in _all_text(document)
    assert "A. A. Objective" not in _all_text(document)


def test_a_table_block_renders_as_a_bordered_table():
    table = {"headers": ["Item", "Cost"], "rows": [["VPN", "NPR 3,390"], ["Licence", "NPR 10,000"]]}
    result = _run(_args(sections=[{"heading": "Cost", "content": [{"table": table}]}]))
    document, _ = _document(result)
    memo_table = document.tables[0]  # the section table comes before the signature grid
    assert [c.text for c in memo_table.rows[0].cells] == ["Item", "Cost"]
    assert [c.text for c in memo_table.rows[2].cells] == ["Licence", "NPR 10,000"]
    assert memo_table.rows[0].cells[0].paragraphs[0].runs[0].bold
