"""Offline tests for the create_pptx local tool (python-pptx -> .pptx).

No network: calls the tool fn directly against a temp-configured fallback file
store and asserts (a) a valid .pptx (zip w/ ppt/presentation.xml) is produced
and stored, (b) deck title / slide titles / bullets / table cells land in the
deck, (c) full Unicode survives, (d) validation returns friendly ERROR strings,
and (e) the slide/bullet caps refuse rather than build.
"""

import asyncio
import zipfile

import pytest

from app.files.store import PPTX_MEDIA_TYPE, file_store
from app.tools.local import pptx as pptx_tool


@pytest.fixture(autouse=True)
def _configure_store(tmp_path):
    file_store.configure(str(tmp_path))
    yield


def _run(args):
    return asyncio.run(pptx_tool.SPEC.func(args))


def _link_id(result: str) -> str:
    assert "Download it at: GET /v1/files/" in result, result
    return result.split("/v1/files/")[1].strip().split()[0]


def _slide_texts(result: str) -> list[str]:
    """Reopen the stored file with python-pptx and return every text run, in order."""
    from pptx import Presentation

    record = file_store.get(_link_id(result))
    assert record is not None
    assert record.media_type == PPTX_MEDIA_TYPE
    with zipfile.ZipFile(record.path) as zf:
        assert "ppt/presentation.xml" in zf.namelist()  # it's a real .pptx
    prs = Presentation(record.path)
    texts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.extend(p.text for p in shape.text_frame.paragraphs)
            if shape.has_table:
                for row in shape.table.rows:
                    texts.extend(cell.text for cell in row.cells)
    return texts


# ---- validation -------------------------------------------------------------


def test_missing_slides_is_error():
    assert _run({}).startswith("ERROR: 'slides' is required")
    assert _run({"slides": []}).startswith("ERROR: 'slides' is required")


def test_slide_not_object_is_error():
    assert _run({"slides": ["just text"]}).startswith("ERROR: slides[0] must be an object")


def test_empty_slide_is_error():
    assert _run({"slides": [{}]}).startswith("ERROR: slides[0] needs at least one of")


def test_bullets_must_be_array_of_strings():
    assert _run({"slides": [{"bullets": "one"}]}).startswith("ERROR: slides[0].bullets must be an array")
    assert _run({"slides": [{"bullets": [1, 2]}]}).startswith("ERROR: slides[0].bullets must be an array")


def test_bad_table_is_error():
    assert _run({"slides": [{"table": "x"}]}).startswith("ERROR: slides[0].table must be an object")
    assert _run({"slides": [{"table": {"rows": []}}]}).startswith("ERROR: slides[0].table.rows must be")
    assert _run({"slides": [{"table": {"rows": ["a"]}}]}).startswith("ERROR: slides[0].table.rows[0] must be")
    assert _run({"slides": [{"table": {"rows": [["a"]], "headers": "h"}}]}).startswith(
        "ERROR: slides[0].table.headers must be"
    )


def test_too_many_slides_is_refused():
    slides = [{"title": f"s{i}"} for i in range(pptx_tool.MAX_SLIDES + 1)]
    result = _run({"slides": slides})
    assert result.startswith("ERROR:") and str(pptx_tool.MAX_SLIDES) in result


def test_too_many_bullets_is_refused():
    bullets = [f"b{i}" for i in range(pptx_tool.MAX_BULLETS_PER_SLIDE + 1)]
    result = _run({"slides": [{"title": "x", "bullets": bullets}]})
    assert result.startswith("ERROR: slides[0].bullets") and str(pptx_tool.MAX_BULLETS_PER_SLIDE) in result


# ---- rendering --------------------------------------------------------------


def test_title_slide_and_bullets_present():
    result = _run(
        {
            "title": "Quarterly Review",
            "subtitle": "Finance",
            "slides": [
                {"title": "Highlights", "bullets": ["Revenue up", "Costs flat"]},
                {"bullets": ["A slide with no title"]},
            ],
        }
    )
    assert "3 slide(s)" not in result  # the count reports CONTENT slides only
    assert "2 slide(s)" in result
    texts = _slide_texts(result)
    assert "Quarterly Review" in texts and "Finance" in texts
    assert "Highlights" in texts
    assert "Revenue up" in texts and "Costs flat" in texts
    assert "A slide with no title" in texts


def test_no_deck_title_means_no_title_slide():
    from pptx import Presentation

    result = _run({"slides": [{"title": "Only"}]})
    record = file_store.get(_link_id(result))
    assert len(Presentation(record.path).slides) == 1


def test_table_slide_present():
    result = _run(
        {
            "slides": [
                {
                    "title": "Numbers",
                    "table": {"headers": ["Month", "Sales"], "rows": [["Jan", 10], ["Feb", 20]]},
                }
            ]
        }
    )
    texts = _slide_texts(result)
    assert "Numbers" in texts
    assert "Month" in texts and "Sales" in texts and "Jan" in texts and "20" in texts


def test_bullets_and_table_on_one_slide():
    result = _run(
        {
            "slides": [
                {
                    "title": "Both",
                    "bullets": ["point one"],
                    "table": {"rows": [["a", "b"]]},
                }
            ]
        }
    )
    texts = _slide_texts(result)
    assert "point one" in texts and "a" in texts and "b" in texts


def test_bullets_and_table_shrink_keeps_body_left_and_width():
    """Regression: shrinking the bullet box for a table must not zero its
    left/width. python-pptx placeholder position setters write a bare xfrm,
    so writing only top/height (without re-asserting left/width) defaults the
    other two to 0 — invisible bullets, even though the text is still in the
    XML. Compare against the layout's own placeholder, the value the shrink
    must preserve."""
    from pptx import Presentation

    result = _run({"slides": [{"title": "Both", "bullets": ["point one"], "table": {"rows": [["a", "b"]]}}]})
    record = file_store.get(_link_id(result))
    prs = Presentation(record.path)
    slide = prs.slides[0]
    body = slide.placeholders[1]
    layout_body = slide.slide_layout.placeholders[1]
    assert body.left == layout_body.left
    assert body.width == layout_body.width


def test_ragged_table_rows_are_padded():
    result = _run({"slides": [{"table": {"headers": ["x", "y", "z"], "rows": [["1"], ["1", "2", "3", "4"]]}}]})
    assert result.startswith("Created"), result
    texts = _slide_texts(result)
    assert "4" in texts  # the widest row sets the column count


def test_full_unicode_preserved():
    result = _run({"title": "Café 🚀 “quotes”", "slides": [{"bullets": ["नेपाल राष्ट्र बैंक", "smile 😀"]}]})
    texts = _slide_texts(result)
    assert "Café 🚀 “quotes”" in texts
    assert "नेपाल राष्ट्र बैंक" in texts and "smile 😀" in texts


def test_filename_suffix_is_forced():
    result = _run({"slides": [{"title": "x"}], "filename": "deck"})
    assert "'deck.pptx'" in result
    result = _run({"slides": [{"title": "x"}]})
    assert "'presentation.pptx'" in result
