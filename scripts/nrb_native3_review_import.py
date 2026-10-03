#!/usr/bin/env python
"""Read the reader's verdicts back, READ-ONLY, and score the pre-registered criterion (spec §6.4).

    .venv/bin/python scripts/nrb_native3_review_import.py \\
        --workbook /home/manoj/nrb-cohort-review.xlsx --out docs/nrb/native3-review.json

PASS = repair-caused word error rate <= 0.5% over the EVIDENCE (cohort) rows,
and no cluster class with wrong words in two rows. Production rows are shipping
QA: reported, never scored. Opened with read_only=True — this script must never
save the workbook (openpyxl would drop the page images).
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path

WER_LIMIT = 0.005


def count_wrong(verdict: str, words: str) -> int | None:
    if verdict == "correct":
        return 0
    if verdict == "wrong words (list them)":
        n = len([w for w in re.split(r"[,\s]+", words or "") if w])
        return n or None  # a verdict with no words listed is unanswered, never clean
    return None


def score(rows: list[dict]) -> dict:
    evidence = [r for r in rows if r["source"] == "cohort" and r["cluster"] not in ("control", "unrepaired")]
    unanswered = [r for r in evidence if count_wrong(r["verdict"], r["wrong"]) is None
                  and r["verdict"] != "wrong throughout"]
    wrong_total, words_total = 0, 0
    wrong_rows_by_class = Counter()
    throughout = 0
    for r in evidence:
        words_total += len(r["repaired"].split())
        n = count_wrong(r["verdict"], r["wrong"])
        if r["verdict"] == "wrong throughout":
            throughout += 1
            wrong_rows_by_class[r["cluster"]] += 1
            continue
        if n:
            wrong_total += n
            wrong_rows_by_class[r["cluster"]] += 1
    wer = round(wrong_total / words_total, 5) if words_total else None
    repeated = sorted(c for c, n in wrong_rows_by_class.items() if n >= 2)
    passed = (not unanswered and throughout == 0 and wer is not None
              and wer <= WER_LIMIT and not repeated)
    return {"evidence_rows": len(evidence), "unanswered": len(unanswered), "words": words_total,
            "wrong_words": wrong_total, "wrong_throughout": throughout, "wer": wer,
            "classes_wrong_twice": repeated, "passed": passed}


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    from openpyxl import load_workbook

    wb = load_workbook(args.workbook, read_only=True, data_only=True)
    ws = wb["native3-repair"]
    it = ws.iter_rows(values_only=True)
    header = next(it)
    col = {name: i for i, name in enumerate(header)}
    rows = []
    for values in it:
        if not values or not values[col["id"]]:
            continue
        rows.append({"id": values[col["id"]], "source": values[col["source"]],
                     "cluster": values[col["cluster"]], "verdict": values[col["verdict"]] or "",
                     "wrong": values[col["wrong words"]] or "",
                     "native_also_wrong": values[col["native also wrong?"]],
                     "repaired": str(values[col["repaired text"]] or "")})
    wb.close()
    result = {"score": score(rows), "rows": rows}
    Path(args.out).write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(result["score"], ensure_ascii=False, indent=2))
    return 0 if result["score"]["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
