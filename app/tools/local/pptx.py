"""Local tool: create_pptx (structured slide deck -> .pptx download link).

Slide-based content model — a deck's unit is a slide, so the model supplies
`slides[]` of {title?, bullets?, table?} plus an optional deck `title`/`subtitle`
that becomes a title slide. The `table` shape is create_docx's exactly, so the
model reuses what it already knows. Rendered with python-pptx on its default
template (full Unicode, no fonts or system libraries). A file tool, so it flows
through the per-user file sink like create_docx/create_pdf.

Caps are module constants, not arguments: the model must not be able to raise
the limit that bounds the file it asks the process to build.
"""

from __future__ import annotations

import asyncio
from io import BytesIO
from typing import Any

from ...files.store import PPTX_MEDIA_TYPE, file_store
from .base import LocalToolSpec

MAX_SLIDES = 50
MAX_BULLETS_PER_SLIDE = 20

_DEFAULT_FILENAME = "presentation.pptx"


def _validate(args: dict[str, Any]) -> tuple[str, str, list[dict], str] | str:
    """Return (title, subtitle, slides, filename) on success, or an ERROR: string."""
    raw_slides = args.get("slides")
    if not isinstance(raw_slides, list) or not raw_slides:
        return "ERROR: 'slides' is required and must be a non-empty array of {title?, bullets?, table?}."
    if len(raw_slides) > MAX_SLIDES:
        return f"ERROR: too many slides ({len(raw_slides)}); the limit is {MAX_SLIDES}. Split the deck."

    for idx, slide in enumerate(raw_slides):
        if not isinstance(slide, dict):
            return f"ERROR: slides[{idx}] must be an object with title/bullets/table."
        bullets = slide.get("bullets")
        table = slide.get("table")
        if not (slide.get("title") or bullets or table is not None):
            return f"ERROR: slides[{idx}] needs at least one of 'title', 'bullets', or 'table'."
        if bullets is not None:
            if not isinstance(bullets, list) or not all(isinstance(b, str) for b in bullets):
                return f"ERROR: slides[{idx}].bullets must be an array of strings."
            if len(bullets) > MAX_BULLETS_PER_SLIDE:
                return (
                    f"ERROR: slides[{idx}].bullets has {len(bullets)} items; "
                    f"the limit is {MAX_BULLETS_PER_SLIDE} per slide. Split across slides."
                )
        if table is not None:
            if not isinstance(table, dict):
                return f"ERROR: slides[{idx}].table must be an object with 'rows' (and optional 'headers')."
            rows = table.get("rows")
            if not isinstance(rows, list) or not rows:
                return f"ERROR: slides[{idx}].table.rows must be a non-empty array of rows."
            for ridx, row in enumerate(rows):
                if not isinstance(row, (list, tuple)):
                    return f"ERROR: slides[{idx}].table.rows[{ridx}] must be an array of cell values."
            headers = table.get("headers")
            if headers is not None and not isinstance(headers, list):
                return f"ERROR: slides[{idx}].table.headers must be an array of column names."

    title = str(args.get("title") or "")
    subtitle = str(args.get("subtitle") or "")
    filename = str(args.get("filename") or _DEFAULT_FILENAME)
    if not filename.lower().endswith(".pptx"):
        filename += ".pptx"
    return title, subtitle, raw_slides, filename


_LAYOUT_TITLE = 0
_LAYOUT_TITLE_AND_CONTENT = 1
_LAYOUT_TITLE_ONLY = 5


def _add_bullets(slide, bullets: list[str]) -> None:
    """Fill the content placeholder (index 1 on the Title-and-Content layout)."""
    body = slide.placeholders[1].text_frame
    body.clear()  # leaves one empty paragraph
    for i, text in enumerate(bullets):
        paragraph = body.paragraphs[0] if i == 0 else body.add_paragraph()
        paragraph.text = text
        paragraph.level = 0


def _add_table(slide, table: dict, top_emu: int) -> None:
    from pptx.util import Emu

    headers = table.get("headers")
    rows = table["rows"]
    ncols = max([len(headers or [])] + [len(r) for r in rows]) or 1
    nrows = len(rows) + (1 if headers else 0)

    prs_width = slide.part.package.presentation_part.presentation.slide_width
    margin = Emu(int(prs_width * 0.05))
    width = Emu(prs_width - 2 * margin)
    row_h = Emu(370_000)  # ~1 cm; python-pptx sizes rows to content anyway
    shape = slide.shapes.add_table(nrows, ncols, margin, Emu(top_emu), width, Emu(row_h * nrows))
    grid = shape.table

    r = 0
    if headers:
        for c in range(ncols):
            grid.cell(0, c).text = str(headers[c]) if c < len(headers) else ""
        r = 1
    for row in rows:
        for c in range(ncols):
            grid.cell(r, c).text = str(row[c]) if c < len(row) else ""
        r += 1


def _build_pptx_bytes(title: str, subtitle: str, slides: list[dict]) -> bytes:
    """Render the deck with python-pptx. Sync — run in a thread."""
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()

    if title:
        ts = prs.slides.add_slide(prs.slide_layouts[_LAYOUT_TITLE])
        ts.shapes.title.text = title
        ts.placeholders[1].text = subtitle  # empty string leaves the placeholder blank

    for spec in slides:
        bullets = spec.get("bullets") or []
        table = spec.get("table")
        layout = _LAYOUT_TITLE_AND_CONTENT if bullets else _LAYOUT_TITLE_ONLY
        slide = prs.slides.add_slide(prs.slide_layouts[layout])
        slide.shapes.title.text = str(spec.get("title") or "")

        table_top = int(prs.slide_height * 0.25)
        if bullets:
            _add_bullets(slide, bullets)
            if table is not None:
                # Shrink the bullet box to the upper part so the table fits below it.
                # Capture the inherited left/width first: python-pptx placeholder
                # position setters write a bare xfrm, so setting only top/height
                # zeroes left/width instead of keeping the layout's values.
                body = slide.placeholders[1]
                left, width = body.left, body.width
                body.top = Emu(int(prs.slide_height * 0.22))
                body.height = Emu(int(prs.slide_height * 0.30))
                body.left = left
                body.width = width
                table_top = int(prs.slide_height * 0.55)
        if table is not None:
            _add_table(slide, table, table_top)

    buffer = BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


async def _create_pptx(args: dict[str, Any]) -> str:
    validated = _validate(args)
    if isinstance(validated, str):  # an ERROR: message
        return validated
    title, subtitle, slides, filename = validated

    try:
        data = await asyncio.to_thread(_build_pptx_bytes, title, subtitle, slides)
    except Exception as exc:  # noqa: BLE001 - report back, don't raise into the loop
        return f"ERROR: failed to build PPTX: {exc}"

    record = await file_store.save(data, filename=filename, media_type=PPTX_MEDIA_TYPE)
    # Same string shape as create_excel/create_pdf/create_docx so the frontend parses it identically.
    return (
        f"Created PowerPoint presentation '{record.filename}' "
        f"({record.size} bytes, {len(slides)} slide(s)). "
        f"Download it at: GET /v1/files/{record.id}"
    )


SPEC = LocalToolSpec(
    name="create_pptx",
    description=(
        "Create a PowerPoint (.pptx) slide deck and return a download link. Use this "
        "when the user asks for a presentation, slides, or a deck. Provide 'slides' "
        "(array of {title?, bullets?, table?}) and optionally a deck 'title' and "
        "'subtitle' (rendered as a title slide) and a 'filename'. Each slide may have "
        "a 'title', 'bullets' (array of short strings, one level) and/or a 'table' "
        "({headers?, rows[][]}). Full Unicode is supported. Keep decks under "
        f"{MAX_SLIDES} slides and {MAX_BULLETS_PER_SLIDE} bullets per slide. For a "
        "document rather than slides use create_docx or create_pdf."
    ),
    parameters={
        "type": "object",
        "properties": {
            "title": {"type": "string", "description": "Optional deck title (makes a title slide)."},
            "subtitle": {"type": "string", "description": "Optional subtitle for the title slide."},
            "slides": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "title": {"type": "string", "description": "Optional slide title."},
                        "bullets": {
                            "type": "array",
                            "items": {"type": "string"},
                            "description": "Optional bullet points, one string each.",
                        },
                        "table": {
                            "type": "object",
                            "properties": {
                                "headers": {
                                    "type": "array",
                                    "items": {"type": "string"},
                                    "description": "Optional column headers.",
                                },
                                "rows": {
                                    "type": "array",
                                    "items": {"type": "array"},
                                    "description": "Rows of cell values; each row is an array.",
                                },
                            },
                            "required": ["rows"],
                            "description": "Optional simple table.",
                        },
                    },
                },
                "description": "Slides in order; each needs at least a title, bullets, or a table.",
            },
            "filename": {
                "type": "string",
                "description": "Output file name, e.g. 'review.pptx' (default 'presentation.pptx').",
            },
        },
        "required": ["slides"],
    },
    func=_create_pptx,
)
