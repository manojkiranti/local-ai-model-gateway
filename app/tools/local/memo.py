"""Local tool: create_memo (structured memo -> branded .docx download link).

Renders the bank's internal memo in ONE fixed format, measured from the
approved memo it reproduces (US Letter, 0.94" margins, Arial): the NIC ASIA
logo top right and **MEMO** top left; a To / From / Subject / Date block ruled
above and below; lettered bold section headings (A., B., …) with
justified body text, bullets and numbered bold sub-headings; and the
"Submitted for review/ support/ approval as proposed" signature grid — three
columns per row of role / blank signing space / name + designation.

The FORMAT is fixed here; everything else is the model's: to, from, subject,
date, every section, and every signatory's role, name and designation. Every
memo has Objective, Background and Recommendation (in that place) and a
signature grid; any other section is optional. The header is laid out with
paragraphs and tab stops, not tables, so no table outline ever shows. Like
create_pptx, the validated args are saved as the file's `preview` so the
frontend can draw the memo without parsing the .docx.

Caps are module constants, not arguments (the create_pptx rule): the model
must not be able to raise the bound on the file it asks the process to build.
"""

from __future__ import annotations

import asyncio
import re
from io import BytesIO
from pathlib import Path
from typing import Any

from ...files.store import DOCX_MEDIA_TYPE, file_store
from ...localtime import today
from .base import LocalToolSpec

_LOGO_PATH = Path(__file__).parent / "assets" / "memo_logo.png"

MAX_SECTIONS = 26  # lettered A-Z
MAX_BLOCKS_PER_SECTION = 40
MAX_BULLETS = 40
MAX_SIGNATORIES = 30

APPROVAL_HEADING = "Submitted for review/ support/ approval as proposed"
MEMO_FONT = "Arial"

_DEFAULT_FILENAME = "memo.docx"
_BOLD = re.compile(r"\*\*(.+?)\*\*")
_ORDINAL = re.compile(r"^(\d{1,2})(st|nd|rd|th)\b(.*)$", re.S)


# --------------------------------------------------------------------------- #
# Validation — tolerant input, one canonical output
# --------------------------------------------------------------------------- #
# The model does not always send the schema's exact shape. Measured on the
# deployed qwen3.5 (2026-09-30): it sent `sections` as a JSON-encoded STRING,
# wrapped blocks as {"type": {"bullets": [...]}}, and wrote sub-bullets as
# "\n  - " lines inside a bullet — then, refused ten times with the same
# message, fell back to create_docx and produced an unbranded document. So the
# input is normalised to ONE canonical block form before anything is checked:
#
#   "text"                                   a paragraph
#   {"heading": "text"}                      a numbered bold sub-heading
#   {"bullets": ["text" | {"text", "bullets": ["text"]}]}
#   {"table": {"headers": [...], "rows": [[...]]}}
#
# and that canonical form is what is rendered and saved as the preview. Every
# refusal carries a working example, so a retry has something to copy.

MAX_TABLE_ROWS = 60

_EXAMPLE = (
    'Example: "sections": [{"heading": "Objective:", "content": ['
    '{"paragraph": "This memo seeks approval for ..."}, '
    '{"bullets": ["First point", "Second point\\n- a sub-point"]}, '
    '{"subheading": "Risk"}, '
    '{"table": {"headers": ["Item", "Cost"], "rows": [["VPN", "NPR 3,390"]]}}]}]'
)

_PARAGRAPH_KEYS = ("paragraph", "text", "body", "para")
_HEADING_KEYS = ("subheading", "heading", "title")
_BULLET_KEYS = ("bullets", "items", "list", "points")
_BULLET_MARK = re.compile(r"^\s*(?:[-*•o▪]|\d+[.)])\s+")


class _Refusal(Exception):
    pass


def _balance_brackets(text: str) -> str:
    """Repair the bracket mistakes the model makes here: a closer left out
    (`["a", "b"}}` for `["a", "b"]}}`), a closer that does not match its
    opener, and closers missing at the end. Scans outside string literals
    only; anything else stays a real error."""
    pairs = {"[": "]", "{": "}"}
    stack: list[str] = []
    out: list[str] = []
    in_string = escaped = False
    for ch in text:
        if in_string:
            out.append(ch)
            if escaped:
                escaped = False
            elif ch == "\\":
                escaped = True
            elif ch == '"':
                in_string = False
            continue
        if ch == '"':
            in_string = True
        elif ch in pairs:
            stack.append(pairs[ch])
        elif ch in "]}":
            if not stack:
                continue  # a stray closer
            if ch != stack[-1] and len(stack) > 1 and ch == stack[-2]:
                out.append(stack.pop())  # a closer was left out: insert it
            ch = stack.pop()
        out.append(ch)
    return "".join(out) + "".join(reversed(stack))


def _decode(value: Any, what: str) -> Any:
    """A JSON-encoded string where an array/object belongs is decoded —
    repairing mismatched brackets if that is all that is wrong with it."""
    if isinstance(value, str):
        text = value.strip()
        if text[:1] in "[{":
            import json

            try:
                return json.loads(text)
            except ValueError:
                pass
            try:
                return json.loads(_balance_brackets(text))
            except ValueError as exc:
                raise _Refusal(
                    f"ERROR: '{what}' was sent as a string that is not valid JSON ({exc}). "
                    f"Send it as a real JSON array, not a string. {_EXAMPLE}"
                ) from None
    return value


def _text(value: Any) -> str | None:
    if isinstance(value, (str, int, float)) and str(value).strip():
        return str(value).strip()
    return None


def _bullet_item(item: Any, where: str) -> str | dict:
    """A bullet: a string (its "- " lines become sub-bullets) or {text, bullets}."""
    if isinstance(item, dict):
        text = next((_text(item.get(k)) for k in ("text", *_PARAGRAPH_KEYS, "title") if _text(item.get(k))), None)
        if text is None:
            raise _Refusal(f"ERROR: {where} needs 'text'. {_EXAMPLE}")
        sub = _decode(item.get("bullets") or item.get("items") or [], where)
        subs = [_BULLET_MARK.sub("", str(x)).strip() for x in (sub if isinstance(sub, list) else [sub])]
        subs = [x for x in subs if x]
        return {"text": text, "bullets": subs} if subs else text
    text = _text(item)
    if text is None:
        raise _Refusal(f"ERROR: {where} must be a non-empty string.")
    lines = [ln for ln in text.splitlines() if ln.strip()]
    head, rest = lines[0].strip(), lines[1:]
    subs = [_BULLET_MARK.sub("", ln).strip() for ln in rest]
    subs = [x for x in subs if x]
    return {"text": _BULLET_MARK.sub("", head), "bullets": subs} if subs else _BULLET_MARK.sub("", head)


def _bullets(items: Any, where: str) -> dict:
    items = _decode(items, where)
    if isinstance(items, str):
        items = [ln for ln in items.splitlines() if ln.strip()]
    if not isinstance(items, list) or not items:
        raise _Refusal(f"ERROR: {where} must be a non-empty array of strings. {_EXAMPLE}")
    if len(items) > MAX_BULLETS:
        raise _Refusal(f"ERROR: {where} has {len(items)} items; the limit is {MAX_BULLETS}.")
    return {"bullets": [_bullet_item(x, f"{where}[{i}]") for i, x in enumerate(items)]}


def _table(table: Any, where: str) -> dict:
    table = _decode(table, where)
    if not isinstance(table, dict):
        raise _Refusal(f"ERROR: {where} must be {{headers?, rows}}. {_EXAMPLE}")
    rows = _decode(table.get("rows"), f"{where}.rows")
    if not isinstance(rows, list) or not rows or not all(isinstance(r, list) for r in rows):
        raise _Refusal(f"ERROR: {where}.rows must be a non-empty array of rows (arrays of cells).")
    if len(rows) > MAX_TABLE_ROWS:
        raise _Refusal(f"ERROR: {where} has {len(rows)} rows; the limit is {MAX_TABLE_ROWS}.")
    headers = table.get("headers")
    out: dict[str, Any] = {"rows": [["" if c is None else str(c) for c in r] for r in rows]}
    if isinstance(headers, list) and headers:
        out["headers"] = ["" if h is None else str(h) for h in headers]
    return out


def _block(block: Any, where: str) -> str | dict | list:
    block = _decode(block, where)
    if isinstance(block, str):
        lines = [ln for ln in block.splitlines() if ln.strip()]
        if not lines:
            raise _Refusal(f"ERROR: {where} is an empty paragraph.")
        if len(lines) > 1 and all(_BULLET_MARK.match(ln) for ln in lines):
            return _bullets(lines, where)
        if len(lines) > 1 and all(_BULLET_MARK.match(ln) for ln in lines[1:]):
            # "Intro:\n- a\n- b" — a lead-in paragraph, then its list.
            return [lines[0].strip(), _bullets(lines[1:], where)]
        return block.strip()
    if not isinstance(block, dict):
        raise _Refusal(f"ERROR: {where} must be an object like {{\"paragraph\": \"...\"}}. {_EXAMPLE}")
    # {"type": {...}} and {"type": "bullets", "items": [...]} wrappers.
    kind = block.get("type")
    if isinstance(kind, dict) and len(block) == 1:
        return _block(kind, where)
    if isinstance(kind, str):
        kind = kind.lower()
        payload = next(
            (block[k] for k in ("content", "value", *_PARAGRAPH_KEYS, *_BULLET_KEYS, *_HEADING_KEYS, "table")
             if k in block),
            None,
        )
        if kind in ("paragraph", "text", "para", "body"):
            return _block(payload, where)
        if kind in ("heading", "subheading", "title"):
            return _block({"subheading": payload}, where)
        if kind in ("bullets", "bullet", "list", "items"):
            return _bullets(payload, f"{where}.bullets")
        if kind == "table":
            return {"table": _table(payload if payload is not None else block, f"{where}.table")}
    if "table" in block:
        return {"table": _table(block["table"], f"{where}.table")}
    for key in _BULLET_KEYS:
        if key in block:
            return _bullets(block[key], f"{where}.{key}")
    for key in _HEADING_KEYS:
        if _text(block.get(key)):
            return {"heading": _text(block[key])}
    for key in _PARAGRAPH_KEYS:
        if key in block:
            return _block(block[key], where)
    raise _Refusal(
        f"ERROR: {where} must have one of 'paragraph', 'subheading', 'bullets' or 'table'. {_EXAMPLE}"
    )


def _normalize_section(sec: Any, where: str) -> dict:
    sec = _decode(sec, where)
    if not isinstance(sec, dict):
        raise _Refusal(f"ERROR: {where} must be an object with 'heading' and 'content'. {_EXAMPLE}")
    heading = _text(sec.get("heading")) or _text(sec.get("title"))
    if heading is None:
        raise _Refusal(f"ERROR: {where} needs a non-empty 'heading'. {_EXAMPLE}")
    heading = re.sub(r"^[A-Z][.)]\s+", "", heading)  # the letter is added here
    content = _decode(sec.get("content", sec.get("blocks")), f"{where}.content")
    if content is None:
        content = []
        for key in ("paragraphs", "body", "text", "paragraph"):
            value = sec.get(key)
            if value:
                content.extend(value if isinstance(value, list) else [value])
        if sec.get("bullets"):
            content.append({"bullets": sec["bullets"]})
        if sec.get("table"):
            content.append({"table": sec["table"]})
    if isinstance(content, (str, dict)):
        content = [content]
    if not isinstance(content, list) or not content:
        raise _Refusal(f"ERROR: {where}.content must be a non-empty array of blocks. {_EXAMPLE}")
    if len(content) > MAX_BLOCKS_PER_SECTION:
        raise _Refusal(
            f"ERROR: {where}.content has {len(content)} blocks; the limit is {MAX_BLOCKS_PER_SECTION}."
        )
    blocks: list = []
    for i, b in enumerate(content):
        out = _block(b, f"{where}.content[{i}]")
        blocks.extend(out if isinstance(out, list) else [out])
    return {"heading": heading, "content": blocks}


# Every memo has these three, in this order around any other sections:
# Objective and Background first, Recommendation last. (key, default heading)
REQUIRED_SECTIONS = (
    ("objective", "Objective:"),
    ("background", "Background:"),
    ("recommendation", "Recommendation and Conclusion:"),
)


def _with_required_sections(args: dict[str, Any], extra: list[dict]) -> list[dict]:
    """Objective, Background, the optional sections, then Recommendation.

    Each required one comes from its own top-level argument or, failing that,
    from the optional `sections` — the model often puts them there — matched
    by heading and moved into place. Missing entirely is a refusal that names
    it, so a memo can never go out without all three."""
    found: dict[str, dict] = {}
    for key, heading in REQUIRED_SECTIONS:
        value = args.get(key)
        if value is not None and value != "" and value != []:
            found[key] = _normalize_section({"heading": heading, "content": value}, key)
            continue
        match = next(
            (s for s in extra if s["heading"].strip().lower().startswith(key)),
            None,
        )
        if match is None:
            raise _Refusal(
                f"ERROR: every memo needs an '{key}' section — pass '{key}' as an array of "
                f"blocks (Objective, Background and Recommendation are required; other "
                f"sections are optional). {_EXAMPLE}"
            )
        extra.remove(match)
        found[key] = match
    return [found["objective"], found["background"], *extra, found["recommendation"]]


def _signatory(sig: Any, where: str) -> dict:
    sig = _decode(sig, where)
    if not isinstance(sig, dict):
        raise _Refusal(f"ERROR: {where} must be {{role, name, designation}}.")
    role, name = _text(sig.get("role")), _text(sig.get("name"))
    if role is None or name is None:
        raise _Refusal(f"ERROR: {where} needs a non-empty 'role' and 'name' ('designation' optional).")
    designation = sig.get("designation", sig.get("title"))
    if designation is not None and not isinstance(designation, str):
        raise _Refusal(f"ERROR: {where}.designation must be a string.")
    return {"role": role, "name": name, "designation": (designation or "").strip()}


def _validate(args: dict[str, Any]) -> tuple[dict[str, Any], str] | str:
    """Return (memo, filename) on success, or an ERROR: string."""
    try:
        for key in ("to", "from", "subject"):
            if _text(args.get(key)) is None:
                return f"ERROR: '{key}' is required and must be a non-empty string."

        sections = _decode(args.get("sections"), "sections")
        if sections is None or sections == "":
            sections = []
        if isinstance(sections, dict):
            sections = [sections]
        if not isinstance(sections, list):
            return f"ERROR: 'sections' must be an array. {_EXAMPLE}"
        extra = [_normalize_section(sec, f"sections[{i}]") for i, sec in enumerate(sections)]
        sections = _with_required_sections(args, extra)
        if len(sections) > MAX_SECTIONS:
            return f"ERROR: the memo has {len(sections)} sections; the limit is {MAX_SECTIONS} (A-Z)."

        signatories = _decode(args.get("signatories"), "signatories")
        if not isinstance(signatories, list) or not signatories:
            return (
                "ERROR: 'signatories' is required: an array of {role, name, designation} "
                "(e.g. role 'Prepared By', 'Supported By', 'Approved By')."
            )
        if len(signatories) > MAX_SIGNATORIES:
            return f"ERROR: 'signatories' has {len(signatories)} items; the limit is {MAX_SIGNATORIES}."
        signatories = [_signatory(sig, f"signatories[{i}]") for i, sig in enumerate(signatories)]
    except _Refusal as refusal:
        return str(refusal)

    date = args.get("date")
    if date is not None and not isinstance(date, str):
        return "ERROR: 'date' must be a string, e.g. '5th August, 2025'."

    memo = {
        "to": _text(args["to"]),
        "from": _text(args["from"]),
        "subject": _text(args["subject"]),
        "date": (date or "").strip() or format_memo_date(),
        "sections": sections,
        "signatories": signatories,
    }
    filename = str(args.get("filename") or _DEFAULT_FILENAME)
    if not filename.lower().endswith(".docx"):
        filename += ".docx"
    return memo, filename


def format_memo_date(d=None) -> str:
    """'5th August, 2025' — the memo's own date style. Server clock, Nepal
    time (app/localtime), never the model's idea of today."""
    d = d or today()
    n = d.day
    suffix = "th" if 11 <= n % 100 <= 13 else {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
    return f"{n}{suffix} {d.strftime('%B')}, {d.year}"


# --------------------------------------------------------------------------- #
# Rendering
# --------------------------------------------------------------------------- #
def _runs(paragraph, text: str, *, size: float, bold: bool = False) -> None:
    """Add `text` to `paragraph`, with **bold** spans in bold."""
    from docx.shared import Pt

    parts = _BOLD.split(text)
    for i, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        run.font.size = Pt(size)
        run.font.name = MEMO_FONT
        run.bold = bold or i % 2 == 1


def _cant_split(row) -> None:
    from docx.oxml import OxmlElement

    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def _set_widths(table, widths) -> None:
    """Fixed column widths. Set on the grid AND every cell, with a fixed
    layout: Word reads the cells, LibreOffice only the grid."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    table.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    for column, width in zip(table.columns, widths):
        column.width = width
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width


def _tight(paragraph, before: float = 0, after: float = 0) -> None:
    from docx.shared import Pt

    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)


def _indent(paragraph, level: int) -> None:
    """Bullet positions measured from the memo: the bullet 0.38" in from the
    margin and its text at 0.63" (level 1); 0.75" / 1.0" (level 2)."""
    from docx.shared import Inches

    fmt = paragraph.paragraph_format
    fmt.left_indent = Inches(0.63 if level == 1 else 1.0)
    fmt.first_line_indent = Inches(-0.25)


def _second_level_bullet_is_o(document) -> None:
    """The memo's sub-bullets are Courier New "o", Word's own second-level
    bullet; the stock template's List Bullet 2 uses "•". Rewrite that one
    level of the numbering definition the style points at."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    style = document.styles["List Bullet 2"].element
    num_pr = style.find(f"{qn('w:pPr')}/{qn('w:numPr')}")
    if num_pr is None:
        return
    num_id = num_pr.find(qn("w:numId")).get(qn("w:val"))
    ilvl_el = num_pr.find(qn("w:ilvl"))
    ilvl = ilvl_el.get(qn("w:val")) if ilvl_el is not None else "0"
    numbering = document.part.numbering_part.element
    num = next(n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == num_id)
    abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
    abstract = next(
        a for a in numbering.findall(qn("w:abstractNum")) if a.get(qn("w:abstractNumId")) == abstract_id
    )
    lvl = next(lv for lv in abstract.findall(qn("w:lvl")) if lv.get(qn("w:ilvl")) == ilvl)
    lvl.find(qn("w:lvlText")).set(qn("w:val"), "o")
    r_pr = lvl.find(qn("w:rPr"))
    if r_pr is None:
        r_pr = OxmlElement("w:rPr")
        lvl.append(r_pr)
    for old in r_pr.findall(qn("w:rFonts")):
        r_pr.remove(old)
    fonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attr), "Courier New")
    r_pr.insert(0, fonts)


def _header(document, content_width) -> None:
    """MEMO on the left and the logo on the right of ONE line — a right tab
    stop, not a table, so no table outline shows in Word or LibreOffice. The
    line is as tall as the logo, so MEMO sits on its bottom edge as in the
    approved memo."""
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Emu, Inches

    p = document.add_paragraph()
    _tight(p, after=4)
    p.paragraph_format.tab_stops.add_tab_stop(Emu(content_width), WD_TAB_ALIGNMENT.RIGHT)
    _runs(p, "MEMO", size=20, bold=True)
    p.add_run("\t")
    if _LOGO_PATH.exists():
        p.add_run().add_picture(str(_LOGO_PATH), width=Inches(1.2))
    else:
        _runs(p, "NIC ASIA", size=16, bold=True)


def _paragraph_rule(paragraph, side: str) -> None:
    """A 1.5pt black rule on one side of a paragraph (w:pBdr)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    edge = OxmlElement(f"w:{side}")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), "12")
    edge.set(qn("w:space"), "1")
    edge.set(qn("w:color"), "000000")
    borders.append(edge)


def _date_runs(paragraph, date: str) -> None:
    """'5th August, 2025' with the ordinal suffix superscripted, as in the memo."""
    from docx.shared import Pt

    m = _ORDINAL.match(date)
    if not m:
        _runs(paragraph, date, size=11)
        return
    _runs(paragraph, m.group(1), size=11)
    sup = paragraph.add_run(m.group(2))
    sup.font.size = Pt(11)
    sup.font.name = MEMO_FONT
    sup.font.superscript = True
    _runs(paragraph, m.group(3), size=11)


def _meta_block(document, memo: dict, content_width) -> None:
    """To / From / Subject / Date as tab-aligned lines, ruled above and below.
    Paragraphs, not a table, so nothing but the two rules is ever drawn; the
    hanging indent keeps a long subject wrapping under its value column."""
    from docx.shared import Inches

    label_w, value_x = Inches(0.82), Inches(1.02)
    rows = [("To", memo["to"]), ("From", memo["from"]), ("Subject", memo["subject"]), ("Date", memo["date"])]
    for i, (label, value) in enumerate(rows):
        p = document.add_paragraph()
        _tight(p, before=3 if i == 0 else 1, after=3 if i == len(rows) - 1 else 1)
        fmt = p.paragraph_format
        fmt.left_indent = value_x
        fmt.first_line_indent = -value_x
        fmt.tab_stops.add_tab_stop(label_w)
        fmt.tab_stops.add_tab_stop(value_x)
        _runs(p, f"{label}\t:\t", size=11, bold=True)
        if label == "Date":
            _date_runs(p, value)
        else:
            _runs(p, value, size=12 if label == "Subject" else 11)
        if i == 0:
            _paragraph_rule(p, "top")
        if i == len(rows) - 1:
            _paragraph_rule(p, "bottom")


def _section(document, letter: str, section: dict) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    p = document.add_paragraph()
    _tight(p, before=14, after=1)
    p.paragraph_format.keep_with_next = True
    run_text = f"{letter}. {section['heading'].strip()}"
    _runs(p, run_text, size=12, bold=True)

    number = 0
    for block in section["content"]:
        if isinstance(block, str):
            para = document.add_paragraph()
            _tight(para, after=6)
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _runs(para, block.strip(), size=11)
        elif "table" in block:
            _memo_table(document, block["table"])
        elif "heading" in block:
            number += 1
            para = document.add_paragraph()
            _tight(para, before=2, after=0)
            para.paragraph_format.keep_with_next = True
            _runs(para, f"{number}. {block['heading'].strip()}", size=11, bold=True)
        else:
            for item in block["bullets"]:
                text = item if isinstance(item, str) else item["text"]
                para = document.add_paragraph(style="List Bullet")
                _tight(para)
                _indent(para, level=1)
                para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                _runs(para, text.strip(), size=11)
                if isinstance(item, dict):
                    for sub in item.get("bullets") or []:
                        sp = document.add_paragraph(style="List Bullet 2")
                        _tight(sp)
                        _indent(sp, level=2)
                        sp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                        _runs(sp, sub.strip(), size=11)


def _memo_table(document, table: dict) -> None:
    """A bordered table inside a section: bold header row, 10pt cells."""
    headers = table.get("headers") or []
    rows = table["rows"]
    ncols = max([len(headers)] + [len(r) for r in rows]) or 1
    t = document.add_table(rows=0, cols=ncols)
    t.style = "Table Grid"
    for values, bold in ([(headers, True)] if headers else []) + [(r, False) for r in rows]:
        cells = t.add_row().cells
        for c in range(ncols):
            para = cells[c].paragraphs[0]
            _tight(para, before=1, after=1)
            _runs(para, str(values[c]) if c < len(values) else "", size=10, bold=bold)
    spacer = document.add_paragraph()
    _tight(spacer, after=2)


def _signatures(document, signatories: list[dict], content_width) -> None:
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Emu, Inches

    p = document.add_paragraph()
    _tight(p, before=18, after=4)
    p.paragraph_format.keep_with_next = True
    _runs(p, APPROVAL_HEADING, size=11, bold=True)
    for run in p.runs:
        run.underline = True

    groups = [signatories[i : i + 3] for i in range(0, len(signatories), 3)]
    table = document.add_table(rows=0, cols=3)
    table.style = "Table Grid"
    col = Emu(content_width // 3)

    for group in groups:
        group = group + [None] * (3 - len(group))
        role_row, sign_row, name_row = table.add_row(), table.add_row(), table.add_row()
        for row in (role_row, sign_row, name_row):
            _cant_split(row)
        sign_row.height = Inches(0.75)
        sign_row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for i, sig in enumerate(group):
            role_p = role_row.cells[i].paragraphs[0]
            sign_p = sign_row.cells[i].paragraphs[0]
            name_p = name_row.cells[i].paragraphs[0]
            for para in (role_p, sign_p, name_p):
                _tight(para)
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            # Keep a signatory's three rows on one page.
            role_p.paragraph_format.keep_with_next = True
            sign_p.paragraph_format.keep_with_next = True
            if sig is None:
                continue
            _runs(role_p, sig["role"], size=10, bold=True)
            _runs(name_p, sig["name"], size=10, bold=True)
            if sig["designation"]:
                name_p.add_run().add_break()
                _runs(name_p, sig["designation"], size=10, bold=True)
    _set_widths(table, [col, col, col])


def _build_memo_bytes(memo: dict) -> bytes:
    """Render the memo with python-docx. Sync — run in a thread."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Inches(0.94))
    section.top_margin, section.bottom_margin = Inches(0.8), Inches(0.9)
    content_width = section.page_width - section.left_margin - section.right_margin

    for name in ("Normal", "List Bullet", "List Bullet 2"):
        style = document.styles[name]
        style.font.name = MEMO_FONT
        style.font.size = Pt(11)
        rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        rfonts.set(qn("w:ascii"), MEMO_FONT)
        rfonts.set(qn("w:hAnsi"), MEMO_FONT)

    _second_level_bullet_is_o(document)
    _header(document, content_width)
    _meta_block(document, memo, content_width)
    for i, sec in enumerate(memo["sections"]):
        _section(document, chr(ord("A") + i), sec)
    _signatures(document, memo["signatories"], content_width)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def logo_bytes() -> bytes | None:
    """The memo's logo, for the frontend preview (GET /v1/branding/memo-logo)."""
    return _LOGO_PATH.read_bytes() if _LOGO_PATH.exists() else None


async def _create_memo(args: dict[str, Any]) -> str:
    validated = _validate(args)
    if isinstance(validated, str):  # an ERROR: message
        return validated
    memo, filename = validated

    try:
        data = await asyncio.to_thread(_build_memo_bytes, memo)
    except Exception as exc:  # noqa: BLE001 - report back, don't raise into the loop
        return f"ERROR: failed to build the memo: {exc}"

    # The validated memo, tagged so the frontend can tell it from a deck.
    record = await file_store.save(
        data, filename=filename, media_type=DOCX_MEDIA_TYPE, preview={"kind": "memo", **memo}
    )
    return (
        f"Created memo '{record.filename}' "
        f"({record.size} bytes, {len(memo['sections'])} section(s), "
        f"{len(memo['signatories'])} signatory(ies)). "
        f"Download it at: GET /v1/files/{record.id}"
    )


_BLOCK_SCHEMA = {
    "type": "object",
    "description": (
        "One block of a section. Use EXACTLY ONE key: 'paragraph' (text), 'subheading' "
        "(a numbered bold sub-heading, 1., 2., …), 'bullets' (array of strings; put "
        "sub-bullets on new lines starting with '- ' inside a bullet string), or 'table'. "
        "Use **text** for bold."
    ),
    "properties": {
        "paragraph": {"type": "string"},
        "subheading": {"type": "string"},
        "bullets": {"type": "array", "items": {"type": "string"}},
        "table": {
            "type": "object",
            "properties": {
                "headers": {"type": "array", "items": {"type": "string"}},
                "rows": {"type": "array", "items": {"type": "array", "items": {"type": "string"}}},
            },
            "required": ["rows"],
        },
    },
}

SPEC = LocalToolSpec(
    name="create_memo",
    description=(
        "Create an official NIC ASIA internal MEMO as a Word (.docx) file in the bank's "
        "fixed memo format (logo, MEMO title, To/From/Subject/Date block, lettered "
        "sections A./B./C., and the 'Submitted for review/ support/ approval as proposed' "
        "signature grid) and return a download link. Use this whenever the user asks for "
        "a memo, approval memo or office note; use create_docx for any other Word document. "
        "Every memo REQUIRES 'to', 'from', 'subject', 'objective', 'background', "
        "'recommendation' and 'signatories'; any other sections are optional and go in "
        "'sections' (placed between Background and Recommendation). 'date' defaults to "
        "today. Pass every array as a real JSON array, never as a string. "
        "Do not put letters (A., B.) in section headings — they are added. "
        "Signatories are shown three per row in the order given, typically one "
        "'Prepared By', then 'Supported By' reviewers, and a final 'Approved By'."
    ),
    parameters={
        "type": "object",
        "properties": {
            "to": {"type": "string", "description": "Recipient, e.g. 'Chief Executive Officer'."},
            "from": {"type": "string", "description": "Sender, e.g. 'Compliance Department'."},
            "subject": {"type": "string", "description": "Memo subject line."},
            "date": {
                "type": "string",
                "description": "Optional, e.g. '5th August, 2025'. Defaults to today.",
            },
            "objective": {
                "type": "array",
                "items": _BLOCK_SCHEMA,
                "description": "REQUIRED. Section A: what the memo seeks (blocks).",
            },
            "background": {
                "type": "array",
                "items": _BLOCK_SCHEMA,
                "description": "REQUIRED. Section B: context and reasons (blocks).",
            },
            "sections": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "heading": {
                            "type": "string",
                            "description": "Heading without its letter, e.g. 'Risk and Mitigation:'.",
                        },
                        "content": {"type": "array", "items": _BLOCK_SCHEMA},
                    },
                    "required": ["heading", "content"],
                },
                "description": (
                    "OPTIONAL extra sections placed between Background and Recommendation, "
                    "e.g. 'Risk and Mitigation:', 'Implementation and Cost:'. Omit if none."
                ),
            },
            "recommendation": {
                "type": "array",
                "items": _BLOCK_SCHEMA,
                "description": "REQUIRED. The last section: what is recommended for approval (blocks).",
            },
            "signatories": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "role": {
                            "type": "string",
                            "description": "e.g. 'Prepared By', 'Supported By', 'Approved By'.",
                        },
                        "name": {"type": "string"},
                        "designation": {"type": "string", "description": "e.g. 'Head Compliance'."},
                    },
                    "required": ["role", "name"],
                },
                "description": "People in the signature grid, in order, three per row.",
            },
            "filename": {
                "type": "string",
                "description": "Output file name, e.g. 'approval-memo.docx' (default 'memo.docx').",
            },
        },
        "required": ["to", "from", "subject", "objective", "background", "recommendation", "signatories"],
    },
    func=_create_memo,
)
